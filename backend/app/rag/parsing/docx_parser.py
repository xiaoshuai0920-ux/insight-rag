"""DOCX parser: reads Heading 1/2/3 and paragraphs, tables -> Markdown tables."""
from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.rag.parsing.base import NormalizedDocument, Section, SourceFormat


class DocxParser:
    supported_format = SourceFormat.DOCX

    def parse(self, file_path: str | Path) -> NormalizedDocument:
        path = Path(file_path)
        doc = DocxDocument(str(path))

        sections: list[Section] = []
        heading_paths: list[str] = []
        current: Section | None = None

        # Iterate body elements in order to keep paragraphs/tables interleaved.
        body = doc.element.body
        para_map = {p._p: p for p in doc.paragraphs}
        table_map = {t._tbl: t for t in doc.tables}

        for child in body.iterchildren():
            if child.tag.endswith("}p"):
                para = para_map.get(child)
                if para is None:
                    continue
                text = para.text.strip()
                if not text:
                    continue
                style_name = (para.style.name or "").lower() if para.style else ""
                level = self._style_level(style_name)
                if level and (style_name.startswith("heading") or style_name.startswith("标题")):
                    heading_paths = heading_paths[: level - 1]
                    heading_paths.append(text)
                    current = Section(
                        heading=text,
                        heading_path=list(heading_paths),
                        content="",
                        level=level,
                    )
                    sections.append(current)
                else:
                    if current is None:
                        current = Section(heading="", content="", level=1)
                        sections.append(current)
                    current.content += text + "\n"
            elif child.tag.endswith("}tbl"):
                table = table_map.get(child)
                if table is not None:
                    md = self._table_to_markdown(table)
                    if md:
                        if current is None:
                            current = Section(heading="", content="", level=1)
                            sections.append(current)
                        current.content += "\n" + md + "\n"

        for sec in sections:
            sec.content = sec.content.strip()

        normalized = NormalizedDocument(
            title=path.stem, format=SourceFormat.DOCX, sections=[s for s in sections if s.content]
        )
        normalized.compute_stats()
        return normalized

    def _style_level(self, style_name: str) -> int | None:
        import re

        m = re.search(r"(?:heading|标题)\s*(\d)", style_name)
        return int(m.group(1)) if m else (1 if style_name in ("heading", "标题") else None)

    def _table_to_markdown(self, table: Table) -> str:
        rows: list[list[str]] = []
        for row in table.rows:
            rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
        rows = [r for r in rows if any(c for c in r)]
        if not rows:
            return ""
        header = rows[0]
        lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * len(header)) + " |"]
        for row in rows[1:]:
            row = (row + [""] * len(header))[: len(header)]
            lines.append("| " + " | ".join(row) + " |")
        return "\n".join(lines)
