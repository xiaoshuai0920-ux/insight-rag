"""Pure-logic unit tests (no DB): parsers, chunking, RRF, BM25 tokenization, evidence."""
from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ.setdefault("DATABASE_URL", "postgresql+psycopg://insightrag:insightrag@localhost:5433/insightrag_test")


# ---------- Markdown parser ----------
class TestMarkdownParser:
    def test_parses_headings_into_sections(self, tmp_path):
        from app.rag.parsing.markdown_parser import MarkdownParser

        f = tmp_path / "doc.md"
        f.write_text(
            "# 差旅管理制度\n\n总则内容。\n\n## 第一章 交通\n\n乘坐高铁。\n\n### 2.1 标准\n\n经济舱。\n",
            encoding="utf-8",
        )
        doc = MarkdownParser().parse(f)
        assert doc.format.value == "markdown"
        assert len(doc.sections) >= 3
        headings = [s.heading for s in doc.sections]
        assert "差旅管理制度" in headings
        assert "第一章 交通" in headings

    def test_strips_front_matter(self, tmp_path):
        from app.rag.parsing.markdown_parser import MarkdownParser

        f = tmp_path / "doc.md"
        f.write_text("---\ntitle: X\n---\n\n# 正文\n\n内容", encoding="utf-8")
        doc = MarkdownParser().parse(f)
        joined = "\n".join(s.content for s in doc.sections)
        headings = "\n".join(s.heading for s in doc.sections)
        assert "title: X" not in joined + headings
        assert "正文" in joined + headings


# ---------- TXT parser ----------
class TestTextParser:
    def test_degrades_to_paragraphs(self, tmp_path):
        from app.rag.parsing.text_parser import TextParser

        f = tmp_path / "a.txt"
        f.write_text("第一段。" * 100 + "\n\n第二段。" * 100, encoding="utf-8")
        doc = TextParser().parse(f)
        assert len(doc.sections) >= 1
        assert all(s.content for s in doc.sections)


# ---------- DOCX parser ----------
class TestDocxParser:
    def test_parses_headings_and_paragraphs(self, tmp_path):
        from docx import Document as DocxDocument

        from app.rag.parsing.docx_parser import DocxParser

        d = DocxDocument()
        d.add_heading("员工手册", level=1)
        d.add_paragraph("欢迎加入公司。")
        d.add_heading("考勤管理", level=2)
        d.add_paragraph("工作时间为九点至六点。")
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "项目"
        table.cell(0, 1).text = "标准"
        table.cell(1, 0).text = "住宿"
        table.cell(1, 1).text = "350元"
        path = tmp_path / "handbook.docx"
        d.save(str(path))

        doc = DocxParser().parse(path)
        headings = [s.heading for s in doc.sections]
        assert "员工手册" in headings
        assert "考勤管理" in headings
        all_text = "\n".join(s.content for s in doc.sections)
        assert "| 项目 | 标准 |" in all_text
        assert "350元" in all_text


# ---------- PDF parser + scanned detection ----------
class TestPdfParser:
    def test_parses_text_pdf_with_pages(self, tmp_path):
        import pymupdf

        from app.rag.parsing.pdf_parser import PdfParser

        pdf = pymupdf.open()
        page = pdf.new_page()
        rect = pymupdf.Rect(50, 50, 545, 750)
        page.insert_textbox(
            rect,
            "第一章 总则\n本制度适用于全体员工，规范日常管理。\n第二章 考勤\n员工每天工作八小时，需要打卡。\n",
            fontsize=12,
            fontname="china-s",
        )
        path = tmp_path / "t.pdf"
        pdf.save(str(path))
        pdf.close()

        doc = PdfParser().parse(path)
        assert doc.page_count == 1
        assert len(doc.sections) >= 1
        all_text = "\n".join((s.content or "") + (s.heading or "") for s in doc.sections)
        assert "总则" in all_text

    def test_scan_pdf_raises(self, tmp_path):
        import pymupdf

        from app.rag.parsing.base import ScanPdfError
        from app.rag.parsing.pdf_parser import PdfParser

        pdf = pymupdf.open()
        pdf.new_page()  # empty page — no text
        path = tmp_path / "scan.pdf"
        pdf.save(str(path))
        pdf.close()

        with pytest.raises(ScanPdfError):
            PdfParser().parse(path)

    def test_unsupported_format(self):
        from app.rag.parsing.base import get_parser

        with pytest.raises(ValueError):
            get_parser(".xyz")


# ---------- Chunking ----------
class TestChunking:
    def _make_doc(self):
        from app.rag.parsing.base import NormalizedDocument, Section

        long_content = "员工出差住宿标准为三百五十元每晚。" * 30
        doc = NormalizedDocument(title="差旅制度", format="markdown")
        doc.sections = [
            Section(heading="第一章 总则", heading_path=["第一章 总则"], content="规范差旅管理。" * 20, level=1),
            Section(heading="第二章 住宿", heading_path=["第一章 总则", "第二章 住宿"], content=long_content, level=1),
            Section(heading="第三章 报销", heading_path=["第三章 报销"], content="五个工作日内提交申请。" * 30, level=1),
        ]
        return doc

    def test_parent_child_structure(self):
        from app.rag.chunking.chunker import build_parent_child_chunks

        parents = build_parent_child_chunks(self._make_doc())
        assert len(parents) >= 3
        for p in parents:
            assert p.children, "每个父块必须至少有一个子块"
            for c in p.children:
                assert c.heading_path
                assert c.content.strip()

    def test_child_respects_token_window(self):
        from app.rag.chunking.chunker import build_parent_child_chunks
        from app.rag.chunking.tokens import estimate_tokens

        parents = build_parent_child_chunks(self._make_doc(), child_tokens=120, child_overlap=10)
        for p in parents:
            for c in p.children:
                assert estimate_tokens(c.content) <= 200  # window + slack

    def test_token_estimator(self):
        from app.rag.chunking.tokens import estimate_tokens

        assert estimate_tokens("") == 0
        assert estimate_tokens("abcdefgh") >= 1
        assert estimate_tokens("一二三四五六七八九十") >= 8


# ---------- RRF ----------
class TestRRF:
    def test_fusion_interleaves_and_dedupes(self):
        from app.rag.retrieval.dense import DenseHit
        from app.rag.retrieval.bm25 import Bm25Hit
        from app.rag.retrieval.rrf import rrf_fuse

        dense = [DenseHit(chunk_id=1, similarity=0.9), DenseHit(chunk_id=2, similarity=0.8), DenseHit(chunk_id=3, similarity=0.7)]
        bm25 = [Bm25Hit(chunk_id=2, score=8.0), Bm25Hit(chunk_id=3, score=7.0), Bm25Hit(chunk_id=4, score=6.0)]
        out = rrf_fuse(dense, bm25, rrf_constant=60, candidate_k=10)
        ids = [c.chunk_id for c in out]
        assert set(ids) == {1, 2, 3, 4}
        # chunk 2 appears in both lists -> should rank first
        assert out[0].chunk_id == 2
        assert out[0].dense_rank == 2 and out[0].bm25_rank == 1
        assert all(c.rrf_rank == i + 1 for i, c in enumerate(out))

    def test_no_alpha_parameter_concept(self):
        # RRF uses only a rank constant; this documents the design decision.
        import inspect

        from app.rag.retrieval import rrf as rrf_mod

        sig = inspect.signature(rrf_mod.rrf_fuse)
        assert "alpha" not in sig.parameters
        assert "rrf_constant" in sig.parameters


# ---------- BM25 tokenization ----------
class TestBm25Tokenizer:
    def test_chinese_tokenization(self):
        from app.rag.retrieval.bm25 import tokenize

        tokens = tokenize("成都出差住宿标准是多少")
        assert "成都" in tokens or "出差" in tokens
        assert "的" not in tokens  # stopwords removed

    def test_query_expansion(self):
        from app.rag.retrieval.bm25 import expand_query_tokens

        expanded = expand_query_tokens("出差报销标准")
        assert "差旅" in expanded  # synonym expansion


# ---------- Evidence ----------
class TestEvidence:
    def test_no_hits_insufficient(self):
        from app.rag.generation.evidence import assess_evidence

        state = assess_evidence([])
        assert state.status == "INSUFFICIENT"

    def test_low_similarity_insufficient(self):
        from app.rag.generation.evidence import assess_evidence
        from app.rag.retrieval.pipeline import RetrievedChunk

        hits = [RetrievedChunk(chunk_id=1, dense_similarity=0.05)]
        state = assess_evidence(hits)
        assert state.status == "INSUFFICIENT"

    def test_good_hits_sufficient(self):
        from app.rag.generation.evidence import assess_evidence
        from app.rag.retrieval.pipeline import RetrievedChunk

        hits = [RetrievedChunk(chunk_id=1, dense_similarity=0.72, content="标准350元")]
        state = assess_evidence(hits)
        assert state.status == "SUFFICIENT"

    def test_value_conflict_conflicting(self):
        from app.rag.generation.evidence import assess_evidence
        from app.rag.retrieval.pipeline import RetrievedChunk

        h1 = RetrievedChunk(chunk_id=1, dense_similarity=0.8, document_title="V3.0", content="二线城市住宿标准为不超过 350 元/晚。")
        h2 = RetrievedChunk(chunk_id=2, dense_similarity=0.79, document_title="V1.5", content="二线城市住宿标准为不超过 280 元/晚。")
        state = assess_evidence([h1, h2])
        assert state.status == "CONFLICTING"
        assert state.conflicts


# ---------- Password policy ----------
class TestPasswordPolicy:
    def test_valid_password(self):
        from app.core.security import validate_password_strength

        assert validate_password_strength("Passw0rd123") is None

    def test_too_short(self):
        from app.core.security import validate_password_strength

        assert validate_password_strength("Ab1") is not None

    def test_needs_letters_and_digits(self):
        from app.core.security import validate_password_strength

        assert validate_password_strength("abcdefgh") is not None
        assert validate_password_strength("12345678") is not None

    def test_jwt_roundtrip(self):
        from app.core.security import create_access_token, decode_access_token

        token = create_access_token(42)
        assert decode_access_token(token) == 42
        assert decode_access_token("garbage") is None
